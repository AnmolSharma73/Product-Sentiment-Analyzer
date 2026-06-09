import io
import json
import base64
import pandas as pd
from typing import Dict, Any, List
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
import docx
from docx.shared import Inches
from openpyxl.styles import PatternFill
from utils.helpers import setup_logger

logger = setup_logger(__name__)

class ReportGenerator:
    """
    Generates downloadable reports in PDF, Excel, DOCX, and JSON formats.
    """
    
    @staticmethod
    def to_json(analysis_df: pd.DataFrame, insights: Dict[str, Any]) -> str:
        """Returns JSON string containing analysis and insights."""
        data = {
            "insights": insights,
            "reviews": analysis_df.to_dict(orient="records")
        }
        return json.dumps(data, indent=2)

    @staticmethod
    def to_excel(analysis_df: pd.DataFrame, insights: Dict[str, Any], product_name: str) -> bytes:
        """Generates an Excel file with multiple sheets."""
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # Sheet 1: Raw Data
            analysis_df.to_excel(writer, sheet_name='Raw Analysis', index=False)
            
            # Sheet 2: Aggregate Summary (simple counts)
            summary_data = []
            aspect_cols = [c for c in analysis_df.columns if c.endswith('_sentiment') and c != 'overall_sentiment']
            for col in aspect_cols:
                aspect = col.replace('aspect_', '').replace('_sentiment', '')
                counts = analysis_df[col].value_counts()
                total = counts.sum()
                if total > 0:
                    summary_data.append({
                        "Aspect": aspect.capitalize(),
                        "Total Mentions": total,
                        "Positive": counts.get("positive", 0),
                        "Negative": counts.get("negative", 0),
                        "Neutral": counts.get("neutral", 0)
                    })
            if summary_data:
                pd.DataFrame(summary_data).to_excel(writer, sheet_name='Aggregate Summary', index=False)
                
            # Sheet 3: AI Insights
            insights_flat = []
            for k, v in insights.items():
                if isinstance(v, list):
                    v = "\n".join([f"- {item}" for item in v])
                elif isinstance(v, dict):
                    v = "\n".join([f"{sub_k}: {sub_v}" for sub_k, sub_v in v.items()])
                insights_flat.append({"Metric": k.replace("_", " ").title(), "Details": str(v)})
                
            pd.DataFrame(insights_flat).to_excel(writer, sheet_name='AI Insights', index=False)
            
            # Apply basic coloring
            workbook = writer.book
            raw_sheet = workbook['Raw Analysis']
            
            pos_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            neg_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            neu_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            
            # Find the overall_sentiment column index
            if 'overall_sentiment' in analysis_df.columns:
                col_idx = analysis_df.columns.get_loc('overall_sentiment') + 1
                for row in range(2, raw_sheet.max_row + 1):
                    val = raw_sheet.cell(row=row, column=col_idx).value
                    if val == 'positive':
                        raw_sheet.cell(row=row, column=col_idx).fill = pos_fill
                    elif val == 'negative':
                        raw_sheet.cell(row=row, column=col_idx).fill = neg_fill
                    elif val == 'neutral':
                        raw_sheet.cell(row=row, column=col_idx).fill = neu_fill

        return output.getvalue()

    @staticmethod
    def to_docx(analysis_df: pd.DataFrame, insights: Dict[str, Any], product_name: str) -> bytes:
        """Generates a DOCX report."""
        doc = docx.Document()
        doc.add_heading(f'Sentiment Analysis Report: {product_name}', 0)
        
        # AI Insights Section
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph(insights.get('executive_summary', 'Not available.'))
        
        doc.add_heading('Top Strengths', level=1)
        for item in insights.get('top_strengths', []):
            doc.add_paragraph(item, style='List Bullet')
            
        doc.add_heading('Critical Issues', level=1)
        for item in insights.get('critical_issues', []):
            doc.add_paragraph(item, style='List Bullet')
            
        doc.add_heading('Actionable Suggestions', level=1)
        for i, item in enumerate(insights.get('actionable_suggestions', []), 1):
            doc.add_paragraph(f"{i}. {item}")
            
        # Sample Data Table
        doc.add_heading('Sample Reviews', level=1)
        sample_df = analysis_df.head(5)
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'ID'
        hdr_cells[1].text = 'Review'
        hdr_cells[2].text = 'Sentiment'
        
        for _, row in sample_df.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row.get('review_id', ''))
            row_cells[1].text = str(row.get('review_text', ''))[:100] + '...'
            row_cells[2].text = str(row.get('overall_sentiment', ''))
            
        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()

    @staticmethod
    def to_pdf(analysis_df: pd.DataFrame, insights: Dict[str, Any], product_name: str, b64_charts: List[str] = None) -> bytes:
        """Generates a PDF report using ReportLab."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []
        
        # Title
        title_style = styles['Title']
        elements.append(Paragraph(f"Sentiment Analysis Report: {product_name}", title_style))
        elements.append(Spacer(1, 12))
        
        # Executive Summary
        h2 = styles['Heading2']
        normal = styles['Normal']
        
        elements.append(Paragraph("Executive Summary", h2))
        elements.append(Paragraph(insights.get('executive_summary', 'Not available.'), normal))
        elements.append(Spacer(1, 12))
        
        # Strengths & Issues
        elements.append(Paragraph("Top Strengths", h2))
        for item in insights.get('top_strengths', []):
            elements.append(Paragraph(f"• {item}", normal))
        elements.append(Spacer(1, 12))
        
        elements.append(Paragraph("Critical Issues", h2))
        for item in insights.get('critical_issues', []):
            elements.append(Paragraph(f"• {item}", normal))
        elements.append(Spacer(1, 12))
        
        # Recommendations
        elements.append(Paragraph("Recommendations", h2))
        for i, item in enumerate(insights.get('actionable_suggestions', []), 1):
            elements.append(Paragraph(f"{i}. {item}", normal))
        elements.append(Spacer(1, 24))
        
        # Insert base64 charts if provided
        if b64_charts:
            elements.append(Paragraph("Visualizations", h2))
            for b64_str in b64_charts:
                try:
                    # Remove header if present
                    if ',' in b64_str:
                        b64_str = b64_str.split(',')[1]
                    img_data = base64.b64decode(b64_str)
                    img_stream = io.BytesIO(img_data)
                    img = Image(img_stream, width=5*Inches, height=3*Inches)
                    elements.append(img)
                    elements.append(Spacer(1, 12))
                except Exception as e:
                    logger.error(f"Failed to embed chart in PDF: {e}")
        
        # Build PDF
        doc.build(elements)
        return output.getvalue()
